"""Post-copy verification and the downloadable report for Site-to-Site jobs.

Verification re-reads every copied item on both sides (batched, 20 per Graph
call) and compares them:

- the destination item must exist;
- files: quickXorHash when both sides expose one, otherwise size.

One expected difference is called out instead of flagged as a failure:
SharePoint writes column values *into* Office files (docx/xlsx/pptx/...)
when metadata is set on them, so their size/hash legitimately changes after
the metadata step — "changed_by_sharepoint".
"""
from __future__ import annotations

import io
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..db import SessionLocal
from ..graph import drive as gd
from ..models import db_models as models

from .errors import NotFound

_OFFICE_EXTENSIONS = {"docx", "docm", "dotx", "xlsx", "xlsm", "xltx", "pptx", "pptm", "potx", "vsdx"}
_COPIED = ("copied", "metadata_done", "metadata_attention")
_CHUNK = 200  # items per progress update


def _ext(name: str) -> str:
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


def _hash(item: dict | None) -> str | None:
    return (((item or {}).get("file") or {}).get("hashes") or {}).get("quickXorHash")


def _compare(row: models.SiteToSiteItem, src: dict | None, dst: dict | None) -> tuple[str, str | None]:
    if dst is None:
        return "missing", "Not found at the destination"
    if "_error" in dst:
        return "error", f"Could not check destination: {dst.get('_message')}"
    if row.kind == "folder":
        return "ok", None
    if src is None:
        return "ok", "Source file no longer exists — destination copy is present"
    if "_error" in src:
        return "error", f"Could not check source: {src.get('_message')}"
    s_size, d_size = src.get("size") or 0, dst.get("size") or 0
    s_hash, d_hash = _hash(src), _hash(dst)
    if s_hash and d_hash and s_hash == d_hash:
        return "ok", None
    if s_size == d_size and not (s_hash and d_hash):
        return "ok", None
    if _ext(row.name) in _OFFICE_EXTENSIONS and row.status in ("metadata_done", "metadata_attention"):
        return "changed_by_sharepoint", (
            f"Office file updated by SharePoint when metadata was applied (source {s_size} B, destination {d_size} B)"
        )
    if s_size != d_size:
        return "size_mismatch", f"Source {s_size} B, destination {d_size} B"
    return "hash_mismatch", "Same size but different content hash"


async def verify_job(job_id: int, run=None) -> dict:
    """Verify every copied item of a job; stores per-item results and a job
    summary. `run` (a CopyRun) gets live verify_done/verify_total updates."""
    with SessionLocal() as db:
        job = db.get(models.SiteToSiteJob, job_id)
        if job is None:
            raise NotFound("Copy job not found")
        job.verify_status = "running"
        db.commit()
        source_drive_id, dest_drive_id = job.source_drive_id, job.dest_drive_id
        rows = db.query(models.SiteToSiteItem).filter_by(job_id=job_id).all()
        for r in rows:
            db.expunge(r)

    copied = [r for r in rows if r.status in _COPIED and r.dest_item_id]
    if run is not None:
        run.verify_total = len(copied)
        run.verify_done = 0
        await run.notify()

    results: dict[int, tuple[str, str | None]] = {}
    try:
        for start in range(0, len(copied), _CHUNK):
            chunk = copied[start : start + _CHUNK]
            files = [r for r in chunk if r.kind == "file"]
            src = await gd.batch_get_items(source_drive_id, [r.source_drive_item_id for r in files], select="id,size,file")
            dst = await gd.batch_get_items(dest_drive_id, [r.dest_item_id for r in chunk], select="id,name,size,file,folder")
            for r in chunk:
                results[r.id] = _compare(r, src.get(r.source_drive_item_id), dst.get(r.dest_item_id))
            if run is not None:
                run.verify_done += len(chunk)
                await run.notify()
    except Exception as e:
        with SessionLocal() as db:
            job = db.get(models.SiteToSiteJob, job_id)
            job.verify_status = "failed"
            job.verify_summary = json.dumps({"error": str(e)})
            db.commit()
        raise

    file_rows = [r for r in rows if r.kind == "file"]
    counts: dict[str, int] = {}
    for status, _ in results.values():
        counts[status] = counts.get(status, 0) + 1
    present_files = [r for r in file_rows if results.get(r.id, ("",))[0] in ("ok", "changed_by_sharepoint")]
    summary = {
        "checked": len(results),
        "ok": counts.get("ok", 0),
        "changed_by_sharepoint": counts.get("changed_by_sharepoint", 0),
        "size_mismatch": counts.get("size_mismatch", 0),
        "hash_mismatch": counts.get("hash_mismatch", 0),
        "missing": counts.get("missing", 0),
        "error": counts.get("error", 0),
        "source_files": len(file_rows),
        "source_folders": sum(1 for r in rows if r.kind == "folder"),
        "source_bytes": sum(r.size or 0 for r in file_rows),
        "dest_files_verified": len(present_files),
        "dest_bytes_verified": sum(r.size or 0 for r in present_files),
        # Skipped on purpose (conflict policy "skip") is not a failure.
        "skipped": sum(1 for r in file_rows if r.status == "skipped"),
        "not_copied": sum(1 for r in file_rows if r.status not in (*_COPIED, "skipped")),
    }
    summary["passed"] = (
        summary["missing"] == 0 and summary["size_mismatch"] == 0 and summary["hash_mismatch"] == 0
        and summary["error"] == 0 and summary["not_copied"] == 0
    )

    with SessionLocal() as db:
        for item_id, (status, detail) in results.items():
            item = db.get(models.SiteToSiteItem, item_id)
            item.verify_status = status
            item.verify_detail = detail
        job = db.get(models.SiteToSiteJob, job_id)
        job.verify_status = "done"
        job.verified_at = datetime.utcnow()
        job.verify_summary = json.dumps(summary)
        db.commit()
    return summary


def build_report(job_id: int) -> tuple[bytes, str]:
    """An .xlsx with a Summary sheet and one row per item (status, verify
    result, errors, metadata / permission notes)."""
    with SessionLocal() as db:
        job = db.get(models.SiteToSiteJob, job_id)
        if job is None:
            raise NotFound("Copy job not found")
        rows = (
            db.query(models.SiteToSiteItem)
            .filter_by(job_id=job_id)
            .order_by(models.SiteToSiteItem.relative_path.asc())
            .all()
        )
        copy_summary = json.loads(job.copy_summary) if job.copy_summary else {}
        verify_summary = json.loads(job.verify_summary) if job.verify_summary else {}

        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        bold = Font(bold=True)
        info = [
            ("Job", f"#{job.id}"),
            ("Source", f"{job.source_site_key} / {job.source_folder_path or '(library root)'}"),
            ("Destination", f"{job.dest_site_key} / {job.dest_folder_path or '(library root)'}"),
            ("Copy status", job.copy_status or "not started"),
            ("Started", job.copy_started_at.isoformat(sep=" ", timespec="seconds") if job.copy_started_at else ""),
            ("Finished", job.copy_finished_at.isoformat(sep=" ", timespec="seconds") if job.copy_finished_at else ""),
            ("Confirmed by", job.confirmed_by_email or ""),
            ("If a file already exists", job.conflict_policy or ""),
            ("Version history", "yes" if job.copy_versions else "no"),
            ("Permissions", "yes" if job.copy_permissions else "no"),
            ("", ""),
            ("Copy", ""),
            *[(f"  {k.replace('_', ' ')}", v) for k, v in copy_summary.items()],
            ("", ""),
            ("Verification", "passed" if verify_summary.get("passed") else ("not run" if not verify_summary else "issues found")),
            *[(f"  {k.replace('_', ' ')}", v) for k, v in verify_summary.items() if k != "passed"],
        ]
        for label, value in info:
            ws.append([label, value])
            if label and not label.startswith("  "):
                ws.cell(ws.max_row, 1).font = bold
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 70

        items_ws = wb.create_sheet("Items")
        headers = ["Path", "Type", "Size (bytes)", "Status", "Verification", "Verification detail", "Error",
                   "Metadata needing attention", "Permissions"]
        items_ws.append(headers)
        for c in items_ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="334155")
        red = PatternFill("solid", fgColor="FEE2E2")
        amber = PatternFill("solid", fgColor="FEF3C7")
        for r in rows:
            meta = json.loads(r.metadata_report) if r.metadata_report else []
            meta_issues = "; ".join(f"{m.get('field')}: {m.get('detail') or m.get('status')}" for m in meta if m.get("status") in ("unmapped", "error"))
            perms = json.loads(r.permissions_report) if r.permissions_report else []
            perm_text = "; ".join(f"{p.get('principal')} [{','.join(p.get('roles') or [])}] {p.get('status')}" for p in perms)
            items_ws.append([r.relative_path, r.kind, r.size or 0, r.status, r.verify_status or "", r.verify_detail or "",
                             r.error or "", meta_issues, perm_text])
            if r.status == "failed" or r.verify_status in ("missing", "size_mismatch", "hash_mismatch", "error"):
                fill = red
            elif r.status in ("skipped", "metadata_attention", "discovered") or r.verify_status == "changed_by_sharepoint":
                fill = amber
            else:
                fill = None
            if fill:
                for c in items_ws[items_ws.max_row]:
                    c.fill = fill
        for col, width in zip("ABCDEFGHI", (70, 8, 14, 18, 22, 50, 50, 50, 50)):
            items_ws.column_dimensions[col].width = width
        items_ws.freeze_panes = "A2"
        items_ws.auto_filter.ref = items_ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"site-to-site-job-{job_id}-report.xlsx"
