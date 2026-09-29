"""REST API for Settings → Tag Configuration (services/tag_config.py).

Mounted from main.py via build_router(require_session). Reads need a
session; every write is admin-only (existing model: ADMIN_EMAILS or a
user_profiles role of Admin — see folder_structure_api._is_admin).

Scope: ?scope=site (default: the session's active site) or ?scope=template
(the default template for new clients; admin only).

  GET    /api/tag-config                     items (+ tree paths), origin, modes
  GET    /api/tag-config/active              Active vocabulary for the frontend/automation
  POST   /api/tag-config/items               create
  PATCH  /api/tag-config/items/{id}          edit / rename / re-parent / status
  POST   /api/tag-config/items/{id}/status   {status}
  DELETE /api/tag-config/items/{id}          hard delete (only unused custom items)
  POST   /api/tag-config/reorder             {level, parent_id, ordered_ids[]}
  GET    /api/tag-config/usage               {item_id: count}
  POST   /api/tag-config/bulk/preview        {mode, rows[] | csv} → colour-coded diff
  POST   /api/tag-config/bulk/commit         {mode, rows[] | csv, confirm, remap}
  POST   /api/tag-config/import/parse        multipart CSV/XLSX → rows[]
  GET    /api/tag-config/export?format=csv|xlsx
  POST   /api/tag-config/reset               {confirm}
  GET    /api/tag-config/snapshots
  POST   /api/tag-config/snapshots/{id}/restore   (id 0 = latest, i.e. "Undo")
  GET/PUT /api/tag-config/modes              per-level Add/Replace selection
  GET    /api/tag-config/vessel-names?refresh=true   Term Store, read-only
  POST   /api/tag-config/domain-sync         {apply, level} Domain/Group/Category column sync
                                              (level defaults to "domain" for back-compat)
  POST   /api/tag-config/retag               {level, old_value, new_value} retag documents still
                                              holding old_value (new_value omitted = clear it).
                                              Runs automatically after a rename or a deactivate;
                                              this is the manual re-run / fix-up-old-data version.
  POST   /api/tag-config/copy-from-site/preview  {source_scope, mode} → diff of copying
                                              another site's (or the template's) Active
                                              tags onto ?scope=
  POST   /api/tag-config/copy-from-site/commit   same body + confirm → applies it
"""
from __future__ import annotations

import csv
import io
import json
import logging

from fastapi import APIRouter, Depends, File, Header, HTTPException, Query, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field

from .config import settings
from .folder_structure_api import _email, _is_admin
from .graph.guard import ProtectedTargetError
from .services import get_backend
from .services import tag_config as tc

log = logging.getLogger(__name__)


class ItemIn(BaseModel):
    level: str
    name: str
    parent_id: int | None = None
    display_name: str | None = None
    folder_name: str | None = None
    code: str | None = None
    description: str | None = None
    sort_order: int | None = None
    status: str | None = "Active"
    attributes: dict | None = None


class ItemPatch(BaseModel):
    name: str | None = None
    display_name: str | None = None
    folder_name: str | None = None
    code: str | None = None
    description: str | None = None
    parent_id: int | None = None
    sort_order: int | None = None
    status: str | None = None
    attributes: dict | None = None


class StatusIn(BaseModel):
    status: str


class ReorderIn(BaseModel):
    level: str
    parent_id: int | None = None
    ordered_ids: list[int]


class BulkRow(BaseModel):
    level: str = Field(alias="Level")
    parent_path: str = Field("", alias="Parent Path")
    name: str = Field(alias="Name")
    code: str | None = Field(None, alias="Code")
    description: str | None = Field(None, alias="Description")
    sort_order: int | None = Field(None, alias="Sort Order")
    status: str | None = Field("Active", alias="Status")

    model_config = {"populate_by_name": True}


class BulkIn(BaseModel):
    mode: str = "add"
    rows: list[BulkRow] = Field(default_factory=list)
    csv: str | None = None
    confirm: bool = False
    remap: dict[int, int] | None = None   # {old parent id: new parent id}
    source: str = "Imported"               # "Custom" for manual Replace lists


class ConfirmIn(BaseModel):
    confirm: bool = False


class ModesIn(BaseModel):
    modes: dict[str, str]


class SyncIn(BaseModel):
    apply: bool = False
    level: str = "domain"   # "domain" | "group" | "category" — see tc.TAXONOMY_SYNC_LEVELS


class RetagIn(BaseModel):
    level: str = "domain"   # "domain" | "group" | "category" — see tc.TAXONOMY_SYNC_LEVELS
    old_value: str
    new_value: str | None = None   # None/omitted = clear the field (deactivate case)


class CopyFromIn(BaseModel):
    source_scope: str
    mode: str = "add"
    confirm: bool = False


def _scope(scope: str | None, *, write: bool, admin: bool) -> str:
    """`scope` is "site" (default — the session's current active site),
    "template" (the default for new clients, admin-only to write), or an
    explicit site_key from the Tag Configuration site picker (GET
    /api/tag-config/sites) — letting an admin manage another configured
    client's Domains/Main Folders/etc. without switching the whole app's
    active site. Any authenticated user can *read* another site's config;
    only an admin can write to it (same rule as "template")."""
    s = (scope or "site").strip().lower()
    if s == "template":
        if write and not admin:
            raise HTTPException(403, "Only administrators can change the default template.")
        return tc.TEMPLATE_SCOPE
    if s == "site":
        return tc.current_site_key()
    if write and not admin:
        raise HTTPException(403, "Only administrators can change another site's tag configuration.")
    return s


def _db():
    if not settings.db_configured:
        raise HTTPException(400, "Tag Configuration needs the database configured.")
    from .db.base import SessionLocal
    return SessionLocal()


def _err(exc: Exception):
    if isinstance(exc, HTTPException):
        raise exc
    if isinstance(exc, tc.TagConfigError):
        detail = {"message": str(exc)}
        if exc.details is not None:
            detail["details"] = exc.details
        raise HTTPException(exc.status, detail)
    if isinstance(exc, ProtectedTargetError):
        raise HTTPException(423, str(exc))
    log.exception("[tag-config] request failed")
    raise HTTPException(500, f"Tag configuration operation failed: {exc}")


async def _audit(action: str, email: str, scope: str, *, level: str | None = None,
                 mode: str | None = None, counts: dict | None = None, target: str = "") -> None:
    """Every change → approval_requests activity row (user, time, level, mode, counts)."""
    backend = get_backend()
    if not hasattr(backend, "_create_activity"):
        return
    c = counts or {}
    msg = (f"{email} {action} tag configuration"
           f"{' (' + tc.LEVEL_LABELS.get(level, level) + ')' if level else ''}"
           f"{' mode=' + mode if mode else ''} scope={scope}"
           + (f": created {c.get('created', 0)}, deactivated {c.get('deactivated', 0)}, "
              f"reactivated {c.get('reactivated', 0)}" if c else "")
           + (f" — {target}" if target else ""))
    try:
        await backend._create_activity(
            action_type=f"tag_config_{action}", requesting_email=email or "unknown",
            requesting_name=(email or "").split("@")[0] or None, department="All Departments",
            target_description=(target or f"Tag configuration ({scope})")[:500],
            payload={"scope": scope, "level": level, "mode": mode}, changes=[c], message=msg,
        )
    except Exception:
        log.exception("[tag-config] audit write failed")


def _sync_domains_later(scope: str, level: str | None) -> None:
    """After a change that can affect a synced level (Domain, Group or
    Category), push the Active items at that level to the library's matching
    SharePoint column in the background (best effort, never blocks). Main
    Folder and Sub Category have no synced column (see TAXONOMY_SYNC_LEVELS)
    and are silently skipped. `level=None` (a bulk import/reset/restore/copy
    that can touch several levels at once) syncs all of them."""
    if scope == tc.TEMPLATE_SCOPE or (level is not None and level not in tc.TAXONOMY_SYNC_LEVELS):
        return
    levels = [level] if level else list(tc.TAXONOMY_SYNC_LEVELS)
    import asyncio

    async def _run():
        for lvl in levels:
            res = await tc.taxonomy_column_sync(tc.get_view(scope, fresh=True), lvl, apply=True)
            if res.get("error"):
                log.warning("[tag-config] %s column sync: %s", lvl, res["error"])
    try:
        asyncio.get_running_loop().create_task(_run())
    except RuntimeError:
        pass


def _retag_later(scope: str, level: str, old_value: str, new_value: str | None) -> None:
    """After a rename (new_value = the new display name) or a deactivate
    (new_value = None) of a Domain/Group/Category item, walk the library in
    the background and fix up documents still holding `old_value` — the
    retro-migration taxonomy_column_sync deliberately never does to the
    column itself. Best-effort: logs, never raises, never blocks the
    request. Skipped for the template scope (no document library of its
    own) and for levels with no synced column."""
    if scope == tc.TEMPLATE_SCOPE or level not in tc.TAXONOMY_SYNC_LEVELS or not (old_value or "").strip():
        return
    import asyncio

    async def _run():
        res = await tc.retag_documents(tc.get_view(scope, fresh=True), level, old_value, new_value)
        if res.get("errors"):
            log.warning("[tag-config] retag %s -> %s: %d/%d updated, errors: %s",
                       old_value, new_value, res.get("updated", 0), res.get("checked", 0), res["errors"][:5])
        else:
            log.info("[tag-config] retag %s -> %s: %d/%d document(s) updated",
                     old_value, new_value, res.get("updated", 0), res.get("checked", 0))
    try:
        asyncio.get_running_loop().create_task(_run())
    except RuntimeError:
        pass


def _rows_from(body: BulkIn) -> list[tc.IncomingRow]:
    if body.csv:
        return tc.parse_import(list(csv.DictReader(io.StringIO(body.csv))))
    return tc.parse_import([{
        "Level": r.level, "Parent Path": r.parent_path, "Name": r.name, "Code": r.code or "",
        "Description": r.description or "", "Sort Order": "" if r.sort_order is None else r.sort_order,
        "Status": r.status or "Active",
    } for r in body.rows])


def _mode_key(scope: str, level: str) -> str:
    return f"tag_config_mode:{scope}:{level}"


def _copy_source_rows(source_key: str) -> list[tc.IncomingRow]:
    """Active items of `source_key`'s tree, in the same shape a bulk import
    file would use, so a copy-from-site reuses the tested Add/Replace planner
    instead of a bespoke merge."""
    src_view = tc.get_view(source_key, fresh=True)
    return tc.parse_import(tc.export_rows(src_view, active_only=True))


def _site_label(key: str) -> str:
    if key == tc.TEMPLATE_SCOPE:
        return "the default template"
    if not settings.db_configured:
        return key
    try:
        from .services.site_provisioning import get_all_configured_sites
        with _db() as db:
            for s in get_all_configured_sites(db):
                if s["site_key"] == key:
                    return s.get("display_name") or key
    except Exception:  # noqa: BLE001
        pass
    return key


def build_router(require_session) -> APIRouter:
    router = APIRouter(prefix="/api/tag-config", tags=["tag-config"])

    def ctx(session=Depends(require_session), x_user_email: str | None = Header(default=None)):
        return {"email": _email(session, x_user_email), "admin": _is_admin(session, x_user_email)}

    def admin_ctx(c=Depends(ctx)):
        if not c["admin"]:
            raise HTTPException(403, "Only administrators can change the tag configuration.")
        return c

    @router.get("/sites")
    async def list_sites(c=Depends(ctx)):
        """Configured clients/sites for the scope picker — lets an admin
        manage one specific site's Tag Configuration without switching the
        whole app's active site. Not the same list as the Sites module's
        provisioning table; hidden sites are excluded."""
        if not settings.db_configured:
            return {"sites": [], "active_site": tc.current_site_key()}
        from .services.site_provisioning import get_all_configured_sites
        with _db() as db:
            sites = get_all_configured_sites(db)
        return {
            "sites": [{"site_key": s["site_key"], "display_name": s.get("display_name") or s["site_key"]}
                      for s in sites],
            "active_site": tc.current_site_key(),
        }

    @router.post("/copy-from-site/preview")
    async def copy_from_site_preview(body: CopyFromIn, scope: str | None = Query(None), c=Depends(admin_ctx)):
        """Diff of copying another configured client's (or the template's)
        Active Domains/Main Folders/Groups/Categories/Sub Categories onto
        `scope`. Fixes a confused "other site" that has no Domain/tags of
        its own by seeding it from a site that is already set up correctly,
        instead of the generic template."""
        try:
            target = _scope(scope, write=False, admin=True)
            source = _scope(body.source_scope, write=False, admin=True)
            if source == target:
                raise HTTPException(400, "Choose a different site to copy from.")
            view = tc.get_view(target, fresh=True)
            with _db() as db:
                use = tc.usage_counts(db, view)
            rows = _copy_source_rows(source)
            if not rows:
                raise HTTPException(400, f"{_site_label(source)} has no Active tag configuration to copy.")
            return {**tc.plan_changes(view, rows, body.mode, use).to_dict(),
                    "source_label": _site_label(source), "target_label": _site_label(target)}
        except ValueError as exc:
            raise HTTPException(400, str(exc))
        except Exception as exc:  # noqa: BLE001
            _err(exc)

    @router.post("/copy-from-site/commit")
    async def copy_from_site_commit(body: CopyFromIn, scope: str | None = Query(None), c=Depends(admin_ctx)):
        if not body.confirm:
            raise HTTPException(400, "Preview the changes and confirm to apply them.")
        try:
            target = _scope(scope, write=True, admin=True)
            source = _scope(body.source_scope, write=False, admin=True)
            if source == target:
                raise HTTPException(400, "Choose a different site to copy from.")
            rows = _copy_source_rows(source)
            if not rows:
                raise HTTPException(400, f"{_site_label(source)} has no Active tag configuration to copy.")
            with _db() as db:
                res = tc.commit_changes(db, target, rows, body.mode, c["email"],
                                        source="Imported", snapshot_mode="Copy from site")
                db.commit()
            tc.invalidate(target)
            await _audit("copy_from_site", c["email"], target, mode=body.mode, counts=res,
                         target=f"copied from {_site_label(source)}")
            _sync_domains_later(target, None)
            return res
        except Exception as exc:  # noqa: BLE001
            _err(exc)

    @router.get("")
    async def list_config(scope: str | None = Query(None), c=Depends(ctx)):
        try:
            key = _scope(scope, write=False, admin=c["admin"])
            view = tc.get_view(key, fresh=True)
            modes = {}
            if settings.db_configured:
                from .db import models
                with _db() as db:
                    for lvl in tc.LEVELS:
                        row = db.query(models.AppSetting).filter_by(key=_mode_key(key, lvl)).one_or_none()
                        modes[lvl] = (row.value if row else None) or "add"
            return {**view.to_public(), "scope": "template" if key == tc.TEMPLATE_SCOPE else "site",
                    "levels": [{"id": l, "label": tc.LEVEL_LABELS[l]} for l in tc.LEVELS],
                    "modes": modes, "is_admin": c["admin"],
                    "folder_name_rules": {"invalid_chars": "".join(sorted(tc.INVALID_CHARS)),
                                          "max_length": tc.MAX_NAME_LEN, "ampersand": "& → and"}}
        except Exception as exc:  # noqa: BLE001
            _err(exc)

    @router.get("/active")
    async def active(_c=Depends(ctx)):
        view = tc.get_view()
        return {
            "site_key": view.site_key, "origin": view.origin,
            "domains": view.domain_names(active_only=True),
            "default_domain": view.default_domain(),
            "domain_aliases": view.domain_alias_map(),
            "tree": [tc.public_row(r, view) for r in view.items(active_only=True)],
        }

    @router.get("/usage")
    async def usage(scope: str | None = Query(None), c=Depends(ctx)):
        try:
            key = _scope(scope, write=False, admin=c["admin"])
            with _db() as db:
                return tc.usage_counts(db, tc.get_view(key, fresh=True))
        except Exception as exc:  # noqa: BLE001
            _err(exc)

    @router.post("/items")
    async def create(body: ItemIn, scope: str | None = Query(None), c=Depends(admin_ctx)):
        try:
            key = _scope(scope, write=True, admin=True)
            with _db() as db:
                row = tc.create_item(db, key, body.model_dump(), c["email"])
                db.commit()
            tc.invalidate(key)
            await _audit("add", c["email"], key, level=body.level, counts={"created": 1}, target=body.name)
            _sync_domains_later(key, body.level)
            return row
        except Exception as exc:  # noqa: BLE001
            _err(exc)

    @router.patch("/items/{item_id}")
    async def patch(item_id: int, body: ItemPatch, scope: str | None = Query(None), c=Depends(admin_ctx)):
        try:
            key = _scope(scope, write=True, admin=True)
            with _db() as db:
                before = tc.item_row(db, key, item_id)
                row = tc.update_item(db, key, item_id, body.model_dump(exclude_unset=True), c["email"])
                db.commit()
            tc.invalidate(key)
            action = "rename" if body.name else "edit"
            await _audit(action, c["email"], key, level=row["level"], target=row["name"])
            _sync_domains_later(key, row["level"])
            if body.name and before["display_name"] != row["display_name"]:
                # Column choices/terms are kept in sync above; this is the
                # separate, heavier step of updating documents that already
                # carry the old value (see tc.retag_documents).
                _retag_later(key, row["level"], before["display_name"], row["display_name"])
            return row
        except Exception as exc:  # noqa: BLE001
            _err(exc)

    @router.post("/items/{item_id}/status")
    async def status(item_id: int, body: StatusIn, scope: str | None = Query(None), c=Depends(admin_ctx)):
        try:
            key = _scope(scope, write=True, admin=True)
            with _db() as db:
                row = tc.set_status(db, key, item_id, body.status, c["email"])
                db.commit()
            tc.invalidate(key)
            action = "activate" if body.status == "Active" else "deactivate"
            await _audit(action, c["email"], key, level=row["level"],
                         counts={"deactivated" if action == "deactivate" else "reactivated": 1}, target=row["name"])
            _sync_domains_later(key, row["level"])
            if action == "deactivate":
                # Column choices no longer offer this value (above); also
                # clear it off documents that already had it, instead of
                # leaving them stuck showing a value nobody can pick anymore.
                _retag_later(key, row["level"], row["display_name"], None)
            return row
        except Exception as exc:  # noqa: BLE001
            _err(exc)

    @router.post("/retag")
    async def retag(body: RetagIn, scope: str | None = Query(None), c=Depends(admin_ctx)):
        """Manual, synchronous retag — for fixing documents that were left
        with a stale value from before this auto-retag existed, or for
        re-running a retag that had errors. Awaited (not fire-and-forget)
        so the admin sees exactly what happened."""
        try:
            key = _scope(scope, write=True, admin=True)
            if body.level not in tc.TAXONOMY_SYNC_LEVELS:
                raise HTTPException(400, f"level must be one of {tuple(tc.TAXONOMY_SYNC_LEVELS)}")
            res = await tc.retag_documents(tc.get_view(key, fresh=True), body.level, body.old_value, body.new_value)
            await _audit("retag", c["email"], key, level=body.level,
                         counts={"updated": res.get("updated", 0)},
                         target=f"{body.old_value} -> {body.new_value or '(cleared)'}")
            return res
        except Exception as exc:  # noqa: BLE001
            _err(exc)

    @router.delete("/items/{item_id}")
    async def delete(item_id: int, scope: str | None = Query(None), c=Depends(admin_ctx)):
        try:
            key = _scope(scope, write=True, admin=True)
            with _db() as db:
                row = tc.delete_item(db, key, item_id, c["email"])
                db.commit()
            tc.invalidate(key)
            await _audit("delete", c["email"], key, level=row["level"], target=row["name"])
            _sync_domains_later(key, row["level"])
            return {"deleted": row}
        except Exception as exc:  # noqa: BLE001
            _err(exc)

    @router.post("/reorder")
    async def do_reorder(body: ReorderIn, scope: str | None = Query(None), c=Depends(admin_ctx)):
        try:
            key = _scope(scope, write=True, admin=True)
            with _db() as db:
                rows = tc.reorder(db, key, body.level, body.parent_id, body.ordered_ids, c["email"])
                db.commit()
            tc.invalidate(key)
            await _audit("reorder", c["email"], key, level=body.level)
            return {"items": rows}
        except Exception as exc:  # noqa: BLE001
            _err(exc)

    @router.post("/bulk/preview")
    async def preview(body: BulkIn, scope: str | None = Query(None), c=Depends(admin_ctx)):
        try:
            key = _scope(scope, write=False, admin=True)
            view = tc.get_view(key, fresh=True)
            with _db() as db:
                use = tc.usage_counts(db, view)
            return tc.plan_changes(view, _rows_from(body), body.mode, use).to_dict()
        except ValueError as exc:
            raise HTTPException(400, str(exc))
        except Exception as exc:  # noqa: BLE001
            _err(exc)

    @router.post("/bulk/commit")
    async def commit(body: BulkIn, scope: str | None = Query(None), c=Depends(admin_ctx)):
        if not body.confirm:
            raise HTTPException(400, "Preview the changes and confirm to apply them.")
        try:
            key = _scope(scope, write=True, admin=True)
            with _db() as db:
                res = tc.commit_changes(db, key, _rows_from(body), body.mode, c["email"],
                                        source=body.source if body.source in tc.SOURCES else "Imported",
                                        remap=body.remap)
                db.commit()
            tc.invalidate(key)
            await _audit("replace" if body.mode == "replace" else "import", c["email"], key,
                         mode=body.mode, counts=res)
            _sync_domains_later(key, None)
            return res
        except Exception as exc:  # noqa: BLE001
            _err(exc)

    @router.post("/import/parse")
    async def parse_file(file: UploadFile = File(...), _c=Depends(admin_ctx)):
        data = await file.read()
        name = (file.filename or "").lower()
        try:
            if name.endswith((".xlsx", ".xlsm")):
                from openpyxl import load_workbook
                ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
                it = ws.iter_rows(values_only=True)
                header = [str(h or "").strip() for h in next(it)]
                records = [dict(zip(header, r)) for r in it]
            else:
                records = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
        except Exception as exc:  # noqa: BLE001
            raise HTTPException(400, f"Could not read the file: {exc}")
        rows = tc.parse_import(records)
        return {"rows": [{"Level": tc.LEVEL_LABELS.get(r.level, r.level), "Parent Path": r.parent_path,
                          "Name": r.name, "Code": r.code or "", "Description": r.description or "",
                          "Sort Order": r.sort_order, "Status": r.status, "line": r.line} for r in rows]}

    @router.get("/export")
    async def export(format: str = Query("csv"), scope: str | None = Query(None), c=Depends(ctx)):
        key = _scope(scope, write=False, admin=c["admin"])
        rows = tc.export_rows(tc.get_view(key, fresh=True))
        cols = ["Level", "Parent Path", "Name", "Code", "Description", "Sort Order", "Status", "Source", "Folder Name"]
        if format == "xlsx":
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Tag Configuration"
            ws.append(cols)
            for r in rows:
                ws.append([r[k] for k in cols])
            buf = io.BytesIO()
            wb.save(buf)
            return Response(buf.getvalue(),
                            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            headers={"Content-Disposition": "attachment; filename=tag-configuration.xlsx"})
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)
        return Response("﻿" + buf.getvalue(), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": "attachment; filename=tag-configuration.csv"})

    @router.post("/reset")
    async def reset(body: ConfirmIn, scope: str | None = Query(None), c=Depends(admin_ctx)):
        if not body.confirm:
            raise HTTPException(400, "Confirm to reset to the default values.")
        try:
            key = _scope(scope, write=True, admin=True)
            with _db() as db:
                res = tc.reset_to_default(db, key, c["email"])
                db.commit()
            tc.invalidate(key)
            await _audit("reset", c["email"], key, counts=res)
            _sync_domains_later(key, None)
            return res
        except Exception as exc:  # noqa: BLE001
            _err(exc)

    @router.get("/snapshots")
    async def snapshots(scope: str | None = Query(None), c=Depends(ctx)):
        from .db import models
        key = _scope(scope, write=False, admin=c["admin"])
        with _db() as db:
            rows = (db.query(models.TagConfigSnapshot).filter_by(site_key=key)
                    .order_by(models.TagConfigSnapshot.id.desc()).limit(50).all())
            return [{"id": r.id, "created_at": r.created_at.isoformat() if r.created_at else None,
                     "changed_by": r.changed_by, "mode": r.mode, "level": r.level,
                     "summary": json.loads(r.summary_json or "{}")} for r in rows]

    @router.post("/snapshots/{snapshot_id}/restore")
    async def restore(snapshot_id: int, scope: str | None = Query(None), c=Depends(admin_ctx)):
        try:
            key = _scope(scope, write=True, admin=True)
            with _db() as db:
                res = tc.restore_snapshot(db, key, snapshot_id or None, c["email"])
                db.commit()
            tc.invalidate(key)
            await _audit("restore", c["email"], key, counts={"restored_snapshot": res["restored_snapshot_id"]})
            _sync_domains_later(key, None)
            return res
        except Exception as exc:  # noqa: BLE001
            _err(exc)

    @router.get("/modes")
    async def get_modes(scope: str | None = Query(None), c=Depends(ctx)):
        from .db import models
        key = _scope(scope, write=False, admin=c["admin"])
        with _db() as db:
            out = {}
            for lvl in tc.LEVELS:
                row = db.query(models.AppSetting).filter_by(key=_mode_key(key, lvl)).one_or_none()
                out[lvl] = (row.value if row else None) or "add"
            return out

    @router.put("/modes")
    async def put_modes(body: ModesIn, scope: str | None = Query(None), c=Depends(admin_ctx)):
        from .db import models
        key = _scope(scope, write=True, admin=True)
        with _db() as db:
            for lvl, mode in body.modes.items():
                if lvl not in tc.LEVELS or mode not in tc.MODES:
                    raise HTTPException(400, f"Invalid mode {lvl}={mode}")
                row = db.query(models.AppSetting).filter_by(key=_mode_key(key, lvl)).one_or_none()
                if row is None:
                    row = models.AppSetting(key=_mode_key(key, lvl))
                    db.add(row)
                row.value, row.updated_by = mode, c["email"]
            db.commit()
        return body.modes

    @router.get("/vessel-names")
    async def vessel_names(refresh: bool = Query(False), scope: str | None = Query(None), c=Depends(ctx)):
        key = _scope(scope, write=False, admin=c["admin"])
        try:
            return await tc.term_store_vessels(site_key=key, refresh=refresh)
        except Exception as exc:  # noqa: BLE001
            _err(exc)

    @router.post("/domain-sync")
    async def domain_sync(body: SyncIn, scope: str | None = Query(None), c=Depends(ctx)):
        """Push a level's Active items onto its real SharePoint column
        (choice list or term set). `body.level` is "domain" (default, for
        back-compat), "group" or "category" — Main Folder and Sub Category
        have no synced column. This is what fixes a site whose SharePoint
        library shows no selectable tags even though Tag Configuration lists
        them: that site's column was never seeded with those choices/terms."""
        level = (body.level or "domain").strip().lower()
        if level not in tc.TAXONOMY_SYNC_LEVELS:
            raise HTTPException(400, f"level must be one of {', '.join(tc.TAXONOMY_SYNC_LEVELS)}")
        if body.apply and not c["admin"]:
            raise HTTPException(403, f"Only administrators can change the SharePoint {tc.LEVEL_LABELS[level]} column.")
        key = _scope(scope, write=body.apply, admin=c["admin"])
        try:
            res = await tc.taxonomy_column_sync(tc.get_view(key, fresh=True), level, apply=body.apply)
            if body.apply and res.get("applied"):
                await _audit(f"{level}_sync", c["email"], key, level=level,
                             counts={"created": len(res.get("add", [])), "deactivated": len(res.get("remove", []))})
            return res
        except Exception as exc:  # noqa: BLE001
            _err(exc)

    return router
