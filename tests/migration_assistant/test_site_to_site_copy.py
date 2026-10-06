"""Site-to-Site copy engine against an in-memory fake of the Graph drive API:
background run, conflict policies, pause/cancel/resume, permissions copy,
verification and the Excel report. No tenant or network needed.

Run: python -m pytest tests/migration_assistant/test_site_to_site_copy.py
"""
import asyncio
import io
import json
import os
import tempfile
import unittest
from unittest import mock

_DB_DIR = tempfile.mkdtemp()
os.environ["MIGRATION_DATABASE_URL"] = f"sqlite:///{_DB_DIR}/s2s_test.db"

from openpyxl import load_workbook  # noqa: E402

from app.migration_assistant.db import SessionLocal, init_db  # noqa: E402
from app.migration_assistant.graph import drive as gd  # noqa: E402
from app.migration_assistant.models import db_models as models  # noqa: E402
from app.migration_assistant.services import (  # noqa: E402
    site_to_site_common as s2s_common,
    site_to_site_mover as mover,
    site_to_site_progress as progress,
    site_to_site_service as service,
    site_to_site_verify as verify,
    term_mapping,
)


class FakeDrives:
    """Two drives ("src", "dst") of folders/files keyed by item id."""

    def __init__(self):
        self.items: dict[str, dict] = {}
        self.parent: dict[str, str] = {}
        self.drive_of: dict[str, str] = {}
        self.perms: dict[str, list] = {}
        self.invites: list[tuple] = []
        self.copy_delay = 0.0
        self._n = 0
        for drive in ("src", "dst"):
            self.items[f"{drive}-root"] = {"id": f"{drive}-root", "name": "", "folder": {}}
            self.drive_of[f"{drive}-root"] = drive

    def _new_id(self, drive):
        self._n += 1
        return f"{drive}-{self._n}"

    def add(self, drive, parent_id, name, *, folder=False, size=0, content="x"):
        iid = self._new_id(drive)
        item = {"id": iid, "name": name, "createdDateTime": f"2026-01-01T00:00:{self._n:02d}Z"}
        if folder:
            item["folder"] = {}
        else:
            item["size"] = size
            item["file"] = {"hashes": {"quickXorHash": f"h-{content}"}}
        self.items[iid] = item
        self.parent[iid] = parent_id
        self.drive_of[iid] = drive
        return iid

    def kids(self, parent_id):
        return [self.items[i] for i, p in self.parent.items() if p == parent_id]

    # --- patched graph.drive functions ---
    async def list_children(self, drive_id, item_id):
        return [dict(c) for c in self.kids(item_id)]

    async def ensure_folder(self, drive_id, parent_id, name):
        for c in self.kids(parent_id):
            if c["name"].lower() == name.lower() and "folder" in c:
                return dict(c)
        return dict(self.items[self.add(drive_id, parent_id, name, folder=True)])

    async def copy_item(self, drive_id, item_id, dest_drive_id, dest_parent_id, new_name, *, conflict_behavior="fail", include_versions=False):
        src = self.items[item_id]
        existing = [c for c in self.kids(dest_parent_id) if c["name"].lower() == new_name.lower()]
        name = new_name
        if existing:
            if conflict_behavior == "fail":
                return "monitor:fail:nameAlreadyExists"
            if conflict_behavior == "replace":
                for c in existing:
                    del self.parent[c["id"]]
            if conflict_behavior == "rename":
                stem, dot, ext = new_name.rpartition(".")
                name = f"{stem} 1{dot}{ext}" if dot else f"{new_name} 1"
        content = src["file"]["hashes"]["quickXorHash"][2:]
        new_id = self.add(dest_drive_id, dest_parent_id, name, size=src["size"], content=content)
        return f"monitor:ok:{new_id}"

    async def poll_copy_status(self, monitor_url, *, timeout_seconds=120):
        if self.copy_delay:
            await asyncio.sleep(self.copy_delay)
        _, state, value = monitor_url.split(":", 2)
        if state == "fail":
            return {"status": "failed", "error": {"code": value}}
        return {"status": "completed", "resourceId": value}

    async def get_item(self, drive_id, item_id):
        if item_id not in self.items or item_id not in self.parent and not item_id.endswith("root"):
            from app.migration_assistant.graph.client import GraphError
            raise GraphError(404, "itemNotFound")
        return dict(self.items[item_id])

    async def batch_get_items(self, drive_id, item_ids, select=""):
        return {i: (dict(self.items[i]) if i in self.items and i in self.parent else None) for i in item_ids}

    async def list_permissions(self, drive_id, item_id):
        return self.perms.get(item_id, [])

    async def invite(self, drive_id, item_id, emails, roles):
        self.invites.append((item_id, tuple(emails), tuple(roles)))
        return {}


class SiteToSiteCopyTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from app.migration_assistant.config import settings
        # These tests wipe the site-to-site tables — never run them against a
        # real database (e.g. if the module's settings were loaded earlier
        # by another test with the default DATABASE_URL).
        if _DB_DIR.replace("\\", "/") not in settings.database_url.replace("\\", "/"):
            raise unittest.SkipTest("Migration Assistant settings already point at a non-test database")
        init_db()

    def setUp(self):
        with SessionLocal() as db:
            db.query(models.SiteToSiteItem).delete()
            db.query(models.SiteToSiteJob).delete()
            db.commit()
        progress._runs.clear()
        self.fake = FakeDrives()
        patches = [
            mock.patch.object(gd, name, getattr(self.fake, name))
            for name in ("list_children", "ensure_folder", "copy_item", "poll_copy_status", "get_item",
                         "batch_get_items", "list_permissions", "invite")
        ]
        patches.append(mock.patch.object(s2s_common, "resolve_site", self._resolve_site))
        patches.append(mock.patch.object(term_mapping, "migrate_item_fields", self._no_metadata))
        patches.append(mock.patch.object(service, "_require_configured", lambda: None))
        for p in patches:
            p.start()
            self.addCleanup(p.stop)

    @staticmethod
    async def _resolve_site(key):
        return {"key": key, "hostname": "contoso.sharepoint.com", "site_path": "sites/dst", "site_id": "site-dst"}

    @staticmethod
    async def _no_metadata(**kwargs):
        return [{"field": "Title", "kind": "text", "status": "applied", "detail": None}]

    def _make_job(self, files: dict[str, int]):
        """Source tree under one selected folder "Docs"; returns job id."""
        f = self.fake
        docs = f.add("src", "src-root", "Docs", folder=True)
        folders = {"": docs}
        with SessionLocal() as db:
            job = models.SiteToSiteJob(status="done", source_site_key="src", source_drive_id="src",
                                       dest_site_key="dst", dest_drive_id="dst", dest_folder_id="dst-root")
            db.add(job)
            db.flush()
            for path, size in files.items():
                parts = path.split("/")
                for depth in range(1, len(parts)):
                    rel = "/".join(parts[:depth])
                    if rel not in folders:
                        parent = folders["/".join(parts[: depth - 1])]
                        folders[rel] = f.add("src", parent, parts[depth - 1], folder=True)
                        db.add(models.SiteToSiteItem(job_id=job.id, source_drive_item_id=folders[rel], kind="folder",
                                                     relative_path=rel, name=parts[depth - 1], status="discovered"))
                iid = f.add("src", folders["/".join(parts[:-1])], parts[-1], size=size, content=path)
                db.add(models.SiteToSiteItem(job_id=job.id, source_drive_item_id=iid, kind="file",
                                             relative_path=path, name=parts[-1], size=size, status="discovered"))
            db.commit()
            return str(job.id)

    def _run(self, coro):
        return asyncio.run(coro)

    def _items(self, job_id):
        with SessionLocal() as db:
            return {i.relative_path: i for i in db.query(models.SiteToSiteItem).filter_by(job_id=int(job_id))}

    def _job(self, job_id):
        with SessionLocal() as db:
            return db.get(models.SiteToSiteJob, int(job_id))

    async def _start_and_wait(self, job_id, **opts):
        await mover.start_copy(job_id, "tester@example.com", **opts)
        await progress.get(int(job_id)).task

    def test_full_copy_runs_in_background_and_verifies(self):
        job_id = self._make_job({"a.pdf": 100, "Sub/b.pdf": 200, "Sub/Deep/c.txt": 50})
        self._run(self._start_and_wait(job_id))
        job = self._job(job_id)
        self.assertEqual(job.copy_status, "completed")
        self.assertEqual(json.loads(job.copy_summary)["files_done"], 3)
        self.assertEqual(json.loads(job.copy_summary)["bytes_done"], 350)
        summary = json.loads(job.verify_summary)
        self.assertTrue(summary["passed"])
        self.assertEqual(summary["ok"], 5)  # 3 files + 2 folders
        dst_names = {i["name"] for i in self.fake.items.values() if self.fake.drive_of[i["id"]] == "dst"}
        self.assertTrue({"a.pdf", "b.pdf", "c.txt", "Sub", "Deep"} <= dst_names)

    def test_conflict_skip_replace_rename(self):
        for policy, expect_status, expect_name in (
            ("skip", "skipped", None), ("replace", "metadata_done", "a.pdf"), ("rename", "metadata_done", "a 1.pdf"),
        ):
            with self.subTest(policy=policy):
                self.setUp()
                self.fake.add("dst", "dst-root", "a.pdf", size=999, content="old")
                job_id = self._make_job({"a.pdf": 100})
                self._run(self._start_and_wait(job_id, conflict_policy=policy))
                item = self._items(job_id)["a.pdf"]
                self.assertEqual(item.status, expect_status)
                self.assertEqual(self._job(job_id).copy_status, "completed")
                names = sorted(c["name"] for c in self.fake.kids("dst-root"))
                if policy == "skip":
                    self.assertEqual(names, ["a.pdf"])
                    self.assertEqual(json.loads(self._job(job_id).verify_summary)["skipped"], 1)
                elif policy == "replace":
                    self.assertEqual(names, ["a.pdf"])
                    self.assertEqual(self.fake.items[item.dest_item_id]["size"], 100)
                else:
                    self.assertEqual(names, ["a 1.pdf", "a.pdf"])
                    self.assertEqual(self.fake.items[item.dest_item_id]["name"], expect_name)

    def test_conflict_fail_marks_item_failed(self):
        self.fake.add("dst", "dst-root", "a.pdf", size=1)
        job_id = self._make_job({"a.pdf": 100, "b.pdf": 5})
        self._run(self._start_and_wait(job_id, conflict_policy="fail"))
        items = self._items(job_id)
        self.assertEqual(items["a.pdf"].status, "failed")
        self.assertEqual(items["b.pdf"].status, "metadata_done")
        self.assertEqual(self._job(job_id).copy_status, "completed_with_errors")

    def test_cancel_then_resume_copies_only_the_rest(self):
        job_id = self._make_job({f"f{i}.pdf": 10 for i in range(12)})
        self.fake.copy_delay = 0.02

        async def scenario():
            await mover.start_copy(job_id, "t@x.com")
            run = progress.get(int(job_id))
            while run.files_done < 4:
                await asyncio.sleep(0.005)
            await mover.cancel(job_id)
            await run.task
            first = run.files_done
            self.assertEqual(self._job(job_id).copy_status, "cancelled")
            self.assertLess(first, 12)
            self.fake.copy_delay = 0
            await mover.resume(job_id, "t@x.com")
            await progress.get(int(job_id)).task
            return first

        self._run(scenario())
        self.assertEqual(self._job(job_id).copy_status, "completed")
        self.assertEqual(len([c for c in self.fake.kids("dst-root") if "file" in c]), 12)  # no duplicates

    def test_pause_blocks_new_files_until_resumed(self):
        job_id = self._make_job({f"f{i}.pdf": 10 for i in range(10)})
        self.fake.copy_delay = 0.02

        async def scenario():
            await mover.start_copy(job_id, "t@x.com")
            run = progress.get(int(job_id))
            while run.files_done < 2:
                await asyncio.sleep(0.005)
            mover.pause(job_id)
            self.assertEqual(self._job(job_id).copy_status, "paused")
            await asyncio.sleep(0.1)  # in-flight files finish
            frozen = run.files_done
            await asyncio.sleep(0.1)
            self.assertEqual(run.files_done, frozen)
            self.assertEqual(run.snapshot()["speed_bps"], 0)
            await mover.resume(job_id, "t@x.com")
            await run.task

        self._run(scenario())
        self.assertEqual(self._job(job_id).copy_status, "completed")

    def test_cannot_start_twice(self):
        job_id = self._make_job({"a.pdf": 1})
        self.fake.copy_delay = 0.05

        async def scenario():
            await mover.start_copy(job_id, "t@x.com")
            with self.assertRaises(mover.Conflict):
                await mover.start_copy(job_id, "t@x.com")
            await progress.get(int(job_id)).task

        self._run(scenario())

    def test_permissions_copy_unique_grants_only(self):
        job_id = self._make_job({"a.pdf": 1})
        src_id = self._items(job_id)["a.pdf"].source_drive_item_id
        self.fake.perms[src_id] = [
            {"roles": ["read"], "inheritedFrom": {"id": "x"}, "grantedToV2": {"user": {"email": "inherited@x.com"}}},
            {"roles": ["write"], "grantedToV2": {"user": {"email": "editor@x.com", "displayName": "Ed"}}},
            {"roles": ["read"], "grantedToV2": {"siteUser": {"loginName": "i:0#.f|membership|reader@x.com"}}},
            {"roles": ["read"], "link": {"scope": "anonymous"}},
            {"roles": ["read"], "grantedToV2": {"siteGroup": {"displayName": "Site Visitors"}}},
        ]
        self._run(self._start_and_wait(job_id, copy_permissions=True))
        granted = {(e[0], r) for _, e, r in self.fake.invites}
        self.assertEqual(granted, {("editor@x.com", ("write",)), ("reader@x.com", ("read",))})
        report = json.loads(self._items(job_id)["a.pdf"].permissions_report)
        self.assertEqual(sorted(p["status"] for p in report), ["applied", "applied", "skipped", "skipped"])

    def test_verification_flags_missing_and_mismatched(self):
        job_id = self._make_job({"a.pdf": 10, "b.pdf": 20, "c.docx": 30})
        self._run(self._start_and_wait(job_id))
        items = self._items(job_id)
        del self.fake.parent[items["a.pdf"].dest_item_id]  # deleted at destination
        self.fake.items[items["b.pdf"].dest_item_id]["size"] = 21
        self.fake.items[items["b.pdf"].dest_item_id]["file"]["hashes"]["quickXorHash"] = "h-other"
        self.fake.items[items["c.docx"].dest_item_id]["size"] = 31  # SharePoint wrote metadata into it
        self.fake.items[items["c.docx"].dest_item_id]["file"]["hashes"]["quickXorHash"] = "h-meta"
        summary = self._run(verify.verify_job(int(job_id)))
        self.assertFalse(summary["passed"])
        items = self._items(job_id)
        self.assertEqual(items["a.pdf"].verify_status, "missing")
        self.assertEqual(items["b.pdf"].verify_status, "size_mismatch")
        self.assertEqual(items["c.docx"].verify_status, "changed_by_sharepoint")

    def test_report_workbook(self):
        job_id = self._make_job({"a.pdf": 10, "Sub/b.pdf": 20})
        self._run(self._start_and_wait(job_id))
        content, filename = verify.build_report(int(job_id))
        self.assertTrue(filename.endswith(".xlsx"))
        wb = load_workbook(io.BytesIO(content))
        self.assertEqual(wb.sheetnames, ["Summary", "Items"])
        paths = [r[0] for r in wb["Items"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(sorted(paths), ["Sub", "Sub/b.pdf", "a.pdf"])

    def test_stream_emits_progress_then_done(self):
        job_id = self._make_job({f"f{i}.pdf": 10 for i in range(5)})
        self.fake.copy_delay = 0.01

        async def scenario():
            await mover.start_copy(job_id, "t@x.com")
            return [json.loads(line) async for line in service.stream_progress(job_id)]

        lines = self._run(scenario())
        self.assertEqual(lines[-1]["type"], "done")
        self.assertEqual(lines[-1]["job"]["copy_status"], "completed")
        progress_lines = [l for l in lines if l["type"] == "progress"]
        self.assertGreaterEqual(len(progress_lines), 2)
        self.assertEqual(progress_lines[-1]["files_done"], 5)
        seen = [e["seq"] for l in progress_lines for e in l["events"]]
        self.assertEqual(len(seen), len(set(seen)))  # every event delivered once


class SiteToSiteRouteTests(SiteToSiteCopyTests):
    """The HTTP surface: confirm returns 202 straight away, the NDJSON stream
    follows the run to the end, controls and the report work over HTTP."""

    def setUp(self):
        super().setUp()
        from fastapi import FastAPI
        from fastapi.testclient import TestClient

        from app.migration_assistant import api

        app = FastAPI()
        api.setup(app, lambda: None)
        self.client = TestClient(app)
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)

    # Only the route tests below run in this class.
    test_full_copy_runs_in_background_and_verifies = None
    test_conflict_skip_replace_rename = None
    test_conflict_fail_marks_item_failed = None
    test_cancel_then_resume_copies_only_the_rest = None
    test_pause_blocks_new_files_until_resumed = None
    test_cannot_start_twice = None
    test_permissions_copy_unique_grants_only = None
    test_verification_flags_missing_and_mismatched = None
    test_report_workbook = None
    test_stream_emits_progress_then_done = None

    def test_confirm_stream_report_over_http(self):
        job_id = self._make_job({"a.pdf": 10, "Sub/b.docx": 20})
        base = "/api/migration-assistant/site-to-site/jobs"
        r = self.client.post(f"{base}/{job_id}/confirm", json={"conflict_policy": "rename", "copy_versions": True},
                             headers={"X-User-Email": "hdr@example.com"})
        self.assertEqual(r.status_code, 202, r.text)
        self.assertEqual(r.json()["copy_status"], "running")
        self.assertEqual(r.json()["conflict_policy"], "rename")

        with self.client.stream("GET", f"{base}/{job_id}/stream") as resp:
            self.assertEqual(resp.headers["content-type"], "application/x-ndjson")
            lines = [json.loads(l) for l in resp.iter_lines() if l]
        self.assertEqual(lines[-1]["type"], "done")
        done_job = lines[-1]["job"]
        self.assertEqual(done_job["copy_status"], "completed")
        self.assertTrue(done_job["verify_summary"]["passed"])
        self.assertEqual(done_job["confirmed_by_email"], "hdr@example.com")

        self.assertEqual(self.client.post(f"{base}/{job_id}/pause").status_code, 409)
        r = self.client.post(f"{base}/{job_id}/confirm", json={"conflict_policy": "bogus"})
        self.assertEqual(r.status_code, 400)

        r = self.client.get(f"{base}/{job_id}/report")
        self.assertEqual(r.status_code, 200)
        self.assertIn("attachment;", r.headers["content-disposition"])
        load_workbook(io.BytesIO(r.content))

        jobs = self.client.get(f"{base}").json()
        self.assertEqual(jobs[0]["id"], job_id)
        self.assertFalse(jobs[0]["live"])

    def test_drives_route_accepts_url_site_keys(self):
        seen = {}

        async def fake_drives(site_key):
            seen["key"] = site_key
            return [{"id": "d1", "name": "Documents"}]

        with mock.patch.object(service, "list_site_drives", fake_drives):
            key = "url:https://contoso.sharepoint.com/sites/Docs"
            r = self.client.get("/api/migration-assistant/site-to-site/drives", params={"site_key": key})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(seen["key"], key)


class SiteUrlParsingTests(unittest.TestCase):
    def test_parse_site_url(self):
        cases = {
            "https://contoso.sharepoint.com/sites/Docs": ("contoso.sharepoint.com", "sites/Docs"),
            "https://contoso.sharepoint.com/sites/Docs/Shared%20Documents/Forms/AllItems.aspx": ("contoso.sharepoint.com", "sites/Docs"),
            "contoso.sharepoint.com/teams/Ops": ("contoso.sharepoint.com", "teams/Ops"),
            "https://Contoso.SharePoint.com/": ("contoso.sharepoint.com", ""),
        }
        for url, expected in cases.items():
            self.assertEqual(s2s_common.parse_site_url(url), expected)

    def test_rejects_non_sharepoint(self):
        with self.assertRaises(s2s_common.BadRequest):
            s2s_common.parse_site_url("https://example.com/sites/x")

    def test_url_key_round_trip(self):
        key = s2s_common.site_key_for_url("https://contoso.sharepoint.com/sites/Docs/Shared Documents")
        self.assertEqual(key, "url:https://contoso.sharepoint.com/sites/Docs")
        site = s2s_common.find_site(key)
        self.assertEqual((site["hostname"], site["site_path"], site["label"]), ("contoso.sharepoint.com", "sites/Docs", "Docs"))


if __name__ == "__main__":
    unittest.main()
